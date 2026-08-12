"""Excel report generation (openpyxl).

The monthly workbook has three sheets:
  Summary  — one row per pump house: inspected date, worker, OK/issue counts
  Issues   — every issue reported in the month, with notes
  Details  — every recorded task result, mirroring the BQ checklist
"""
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import checklists, db

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ISSUE_FILL = PatternFill("solid", fgColor="FCE4EC")
OK_FILL = PatternFill("solid", fgColor="E8F5E9")
MISSING_FILL = PatternFill("solid", fgColor="FFF3E0")

RESULT_LABEL = {"ok": "OK", "issue": "ISSUE", "skipped": "N/A"}


def _style_header(ws, ncols: int):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def _autosize(ws, widths: dict[int, int]):
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def monthly_report(conn, year: int, month: int) -> Path:
    inspections = db.submitted_in_month(conn, year, month)
    insp_by_house: dict[str, list] = {}
    for i in inspections:
        insp_by_house.setdefault(i["pump_house"], []).append(i)
    items = db.items_for_inspections(conn, [i["id"] for i in inspections])
    insp_by_id = {i["id"]: i for i in inspections}

    wb = Workbook()

    # ------------------------------------------------------------ Summary
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Pump House", "Name", "Inspected", "Date", "Worker",
               "Tasks OK", "Issues", "Skipped"])
    for house in db.list_pump_houses(conn):
        house_insps = insp_by_house.get(house["code"], [])
        if house_insps:
            ids = {i["id"] for i in house_insps}
            house_items = [it for it in items if it["inspection_id"] in ids]
            ok = sum(1 for it in house_items if it["result"] == "ok")
            issue = sum(1 for it in house_items if it["result"] == "issue")
            skip = sum(1 for it in house_items if it["result"] == "skipped")
            last = house_insps[-1]
            row = [house["code"], house["name"], "YES",
                   last["submitted_at"][:10], last["worker_name"],
                   ok, issue, skip]
            ws.append(row)
            fill = ISSUE_FILL if issue else OK_FILL
        else:
            ws.append([house["code"], house["name"], "NO", "", "", "", "", ""])
            fill = MISSING_FILL
        for col in range(1, 9):
            ws.cell(row=ws.max_row, column=col).fill = fill
    _style_header(ws, 8)
    _autosize(ws, {1: 10, 2: 30, 3: 10, 4: 12, 5: 22, 6: 10, 7: 8, 8: 9})

    # ------------------------------------------------------------- Issues
    ws = wb.create_sheet("Issues")
    ws.append(["Date", "Pump House", "Name", "Category", "Task", "Value",
               "Note", "Photo", "Worker"])
    for it in items:
        if it["result"] != "issue":
            continue
        insp = insp_by_id[it["inspection_id"]]
        task = checklists.task_of(it["category"], it["task_id"])
        cat = checklists.CATEGORY_BY_KEY[it["category"]]
        ws.append([insp["submitted_at"][:10], insp["pump_house"],
                   insp["pump_house_name"], cat["name"], task["desc"],
                   it["value"] or "", it["note"] or "",
                   "YES" if it["photo_file_id"] else "",
                   insp["worker_name"]])
    _style_header(ws, 9)
    _autosize(ws, {1: 12, 2: 11, 3: 26, 4: 26, 5: 45, 6: 14, 7: 40, 8: 7, 9: 22})

    # ------------------------------------------------------------ Details
    ws = wb.create_sheet("Details")
    ws.append(["Date", "Pump House", "Category", "#", "Task", "Frequency",
               "Result", "Value", "Note", "Worker"])
    for it in items:
        insp = insp_by_id[it["inspection_id"]]
        task = checklists.task_of(it["category"], it["task_id"])
        cat = checklists.CATEGORY_BY_KEY[it["category"]]
        ws.append([insp["submitted_at"][:10], insp["pump_house"], cat["name"],
                   it["task_id"], task["desc"],
                   checklists.FREQ_LABELS[it["freq"]],
                   RESULT_LABEL.get(it["result"], ""), it["value"] or "",
                   it["note"] or "", insp["worker_name"]])
        if it["result"] == "issue":
            for col in range(1, 11):
                ws.cell(row=ws.max_row, column=col).fill = ISSUE_FILL
    _style_header(ws, 10)
    _autosize(ws, {1: 12, 2: 11, 3: 26, 4: 4, 5: 50, 6: 14, 7: 8, 8: 14,
                   9: 35, 10: 22})

    out = Path(tempfile.gettempdir()) / f"BPPM_report_{year:04d}-{month:02d}.xlsx"
    wb.save(out)
    return out
